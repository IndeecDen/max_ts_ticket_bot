"""Standalone bot XLSX exporter: a snapshot of the shared statistics result."""
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.roles import ManagementError

DAY_MS = 86400000
TIME_FORMAT = '[h]:mm:ss'


def export_statistics(report, output):
    path = Path(output).resolve()
    if path.suffix.lower() != '.xlsx':
        raise ManagementError('Для Excel-выгрузки укажите файл с расширением .xlsx.')
    if path.exists():
        raise ManagementError('Файл уже существует. Укажите новое имя: перезапись отчётов отключена.')
    groups = report['groups']
    if len(groups) > 1048560:
        raise ManagementError('Слишком много строк для одного листа Excel; сократите период.')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Статистика'
    sheet.sheet_view.showGridLines = False
    book.properties.creator = 'MAX ticket bot'
    book.properties.title = 'Статистика завершённых заявок'

    def put(row, column, value, number_format=None):
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            # IDs exceed Excel's 15-digit numeric precision; all text stays literal.
            cell.data_type = 's'
        cell.font = Font(name='Calibri', size=11, color='203047')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if number_format:
            cell.number_format = number_format
        return cell

    def banner(row, text, fill=None, size=11):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = put(row, 1, text)
        cell.font = Font(name='Calibri', size=size, bold=fill is not None,
                         color='FFFFFF' if fill else '203047')
        if fill:
            for column in range(1, 10):
                sheet.cell(row, column).fill = PatternFill('solid', fgColor=fill)

    period, totals = report['period'], report['totals']
    banner(1, 'MAX · Статистика завершённых заявок', '17365D', 18)
    sheet.row_dimensions[1].height = 36
    banner(2, f"Период завершения: {period['start_date']} — {period['end_date']} включительно")
    scope = 'Все исполнители' if report['specialist_id'] is None else f"Исполнитель: {report['specialist_id']}"
    banner(3, f"{scope} · Часовой пояс: {period['timezone']}")
    for row in (2, 3):
        sheet.row_dimensions[row].height = 24
    for row, label, value, fmt in [
        (5, 'Закрыто заявок', totals['closed'], '0'),
        (6, 'Среднее ожидание', totals['wait_ms'] / totals['wait_samples'] / DAY_MS if totals['wait_samples'] else None, TIME_FORMAT),
        (7, 'Среднее выполнение', totals['work_ms'] / totals['work_samples'] / DAY_MS if totals['work_samples'] else None, TIME_FORMAT),
        (8, 'Измерений ожидания', totals['wait_samples'], '0'),
        (9, 'Измерений выполнения', totals['work_samples'], '0'),
    ]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        put(row, 1, label)
        put(row, 3, value if value is not None else 'Нет данных', fmt if value is not None else None)
        sheet.row_dimensions[row].height = 23
    banner(11, 'Ожидание: от первого сообщения до взятия. Выполнение: от первого взятия до завершения.')
    banner(12, 'После переназначения заявка относится к последнему исполнителю. Некорректные интервалы исключены из средних.')
    banner(13, 'Снимок данных: значения не пересчитываются при изменении ячеек. Пустая длительность означает отсутствие измерений.')
    for row in (11, 12, 13):
        sheet.row_dimensions[row].height = 30
    headers = ['MAX ID чата', 'MAX ID исполнителя', 'Закрыто', 'Ожидание, сумма',
               'Измерений ожидания', 'Ожидание, среднее', 'Выполнение, сумма',
               'Измерений выполнения', 'Выполнение, среднее']
    for column, header in enumerate(headers, 1):
        cell = put(16, column, header)
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    sheet.row_dimensions[16].height = 42
    for row, group in enumerate(groups, 17):
        wait_count, work_count = group['wait_samples'], group['work_samples']
        values = [str(group['chat_id']), str(group['specialist_id']) if group['specialist_id'] is not None else 'Не указан',
                  group['closed'], group['wait_ms'] / DAY_MS if wait_count else None,
                  wait_count, group['wait_ms'] / wait_count / DAY_MS if wait_count else None,
                  group['work_ms'] / DAY_MS if work_count else None,
                  work_count, group['work_ms'] / work_count / DAY_MS if work_count else None]
        for column, value in enumerate(values, 1):
            cell = put(row, column, value, '@' if column in (1, 2) else TIME_FORMAT if column in (4, 6, 7, 9) else '0')
            if row % 2:
                cell.fill = PatternFill('solid', fgColor='EDF3F8')
        sheet.row_dimensions[row].height = 28
    if not groups:
        banner(17, 'Нет завершённых заявок за выбранный период.')
        sheet.row_dimensions[17].height = 28
    for column, width in enumerate((25, 25, 13, 20, 19, 20, 20, 21, 21), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = 'C17'
    if groups:
        sheet.auto_filter.ref = f'A16:I{16 + len(groups)}'
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '16:16'
    sheet.print_area = f'A1:I{max(17, 16 + len(groups))}'
    sheet.oddFooter.center.text = 'Страница &P из &N'
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix='.max-report-', suffix='.xlsx', delete=False) as temp:
            temporary = Path(temp.name)
        book.save(temporary)
        # Atomic publication on a local filesystem, refusing even a racing writer.
        try:
            os.link(temporary, path)
        except FileExistsError:
            raise ManagementError('Файл уже существует. Укажите новое имя.') from None
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return path
