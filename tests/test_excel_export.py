import asyncio
import copy
import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from app.config import load_settings
from app.main import show_statistics
from app.services.excel_export import export_statistics, DAY_MS
from app.services.roles import ManagementError


def sample_report():
    totals = {'closed': 3, 'wait_ms': 120000, 'wait_samples': 2,
              'work_ms': 90061000, 'work_samples': 2}
    return {'period': {'start_date': '2026-09-01', 'end_date': '2026-09-22', 'timezone': 'Europe/Moscow'},
            'specialist_id': None, 'totals': totals,
            'groups': [{'chat_id': -9223372036854775807, 'specialist_id': 9223372036854775806, **totals}]}


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root / 'Отчёты' / 'Статистика.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def test_roundtrip_ids_counts_and_durations_over_24_hours(self):
        export_statistics(sample_report(), self.path)
        book = load_workbook(self.path, data_only=True)
        try:
            sheet = book['Статистика']
            self.assertEqual(sheet['A17'].value, '-9223372036854775807')
            self.assertEqual(sheet['B17'].value, '9223372036854775806')
            self.assertEqual(sheet['A17'].data_type, 's')
            self.assertEqual(sheet['C5'].value, 3)
            self.assertEqual(sheet['C17'].value, 3)
            self.assertEqual(sheet['G17'].value.total_seconds(), 90061)
            self.assertEqual(sheet['G17'].number_format, '[h]:mm:ss')
            self.assertEqual(sheet.freeze_panes, 'C17')
            self.assertEqual(sheet.auto_filter.ref, 'A16:I17')
        finally:
            book.close()

    def test_empty_report_and_missing_samples(self):
        report = sample_report()
        report['groups'] = []
        report['totals'] = {key: 0 for key in report['totals']}
        export_statistics(report, self.path)
        book = load_workbook(self.path)
        try:
            sheet = book.active
            self.assertEqual(sheet['C5'].value, 0)
            self.assertEqual(sheet['C6'].value, 'Нет данных')
            self.assertIn('Нет завершённых', sheet['A17'].value)
            self.assertIsNone(sheet.auto_filter.ref)
        finally:
            book.close()

    def test_export_is_snapshot_and_text_cannot_become_formula(self):
        report = sample_report()
        report['groups'][0]['specialist_id'] = '=HYPERLINK("https://example.com")'
        original = copy.deepcopy(report)
        export_statistics(report, self.path)
        self.assertEqual(report, original)
        book = load_workbook(self.path, data_only=False)
        try:
            self.assertEqual(book.active['B17'].data_type, 's')
            self.assertFalse(any(c.data_type == 'f' for row in book.active for c in row))
        finally:
            book.close()

    def test_existing_file_is_never_overwritten(self):
        self.path.parent.mkdir()
        self.path.write_bytes(b'existing user file')
        with self.assertRaises(ManagementError):
            export_statistics(sample_report(), self.path)
        self.assertEqual(self.path.read_bytes(), b'existing user file')

    def test_atomic_publish_race_keeps_competing_file_and_removes_temp(self):
        def competitor(source, destination):
            Path(destination).write_bytes(b'other writer')
            raise FileExistsError()
        with patch('app.services.excel_export.os.link', side_effect=competitor):
            with self.assertRaises(ManagementError):
                export_statistics(sample_report(), self.path)
        self.assertEqual(self.path.read_bytes(), b'other writer')
        self.assertEqual(list(self.path.parent.glob('.max-report-*')), [])

    def test_failed_save_has_no_partial_final_file(self):
        with patch('app.services.excel_export.Workbook.save', side_effect=OSError('disk')):
            with self.assertRaises(OSError):
                export_statistics(sample_report(), self.path)
        self.assertFalse(self.path.exists())
        self.assertEqual(list(self.path.parent.glob('.max-report-*')), [])
        with self.assertRaises(ManagementError):
            export_statistics(sample_report(), self.root / 'report.csv')

    def test_cli_export_uses_statistics_result_and_project_relative_path(self):
        args = SimpleNamespace(xlsx='reports/test.xlsx', json_output=False, from_date=None,
                               to_date=None, period='week', user_id=99)
        settings = load_settings(self.root, {})
        with patch('app.main.PROJECT_ROOT', self.root), \
             patch('app.main.Statistics.read', return_value=sample_report()) as read, \
             redirect_stdout(io.StringIO()):
            self.assertEqual(asyncio.run(show_statistics(settings, args)), 0)
        read.assert_awaited_once_with(['week'], 'Europe/Moscow', 99)
        self.assertTrue((self.root / 'reports/test.xlsx').is_file())
        args.json_output = True
        with self.assertRaises(ManagementError):
            asyncio.run(show_statistics(settings, args))
